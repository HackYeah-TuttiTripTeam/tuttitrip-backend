"""Read the cells of ``miasta.xlsx`` (the Google Sheet's XLSX export)."""

from collections.abc import Iterable
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from tuttitrip.places.logic.sheet_rows import (
    Cell,
    RowError,
    Sheet,
    SheetImportError,
    SheetRow,
)


def read_workbook(path: Path) -> dict[str, Sheet]:
    """Load every worksheet: the first row is the header, values only.

    Formulas are read as the values Google exported; formatting is ignored, so
    the parser depends on headers, not on how the cells look.

    Args:
        path: The XLSX file.

    Returns:
        The worksheets by name; rows with no value at all are dropped.

    Raises:
        SheetImportError: When the file is not a readable XLSX workbook.
    """
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise SheetImportError(
            [RowError(path.name, 0, f"to nie jest XLSX: {exc}")]
        ) from exc
    try:
        return {
            sheet.title: _read_sheet(sheet.title, sheet.iter_rows(values_only=True))
            for sheet in workbook
        }
    finally:
        workbook.close()


def _cell(value: object) -> Cell:
    if value is None or isinstance(value, str | int | float | bool | date):
        return value
    return float(value) if isinstance(value, Decimal) else str(value)


def _read_sheet(name: str, worksheet: Iterable[tuple[object, ...]]) -> Sheet:
    """Split a worksheet into its header and data rows.

    Args:
        name: Sheet title.
        worksheet: Rows of cell values.

    Returns:
        The sheet; an empty one when the worksheet has no rows.
    """
    rows = iter(worksheet)
    header = next(rows, ())
    columns = tuple("" if cell is None else str(cell).strip() for cell in header)
    data = [
        SheetRow(
            number, {c: _cell(v) for c, v in zip(columns, values, strict=False) if c}
        )
        for number, values in enumerate(rows, start=2)
        if any(v is not None and str(v).strip() for v in values)
    ]
    return Sheet(name, columns, tuple(data))
